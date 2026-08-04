"""A native PowerPoint table of Channel-1's mROI error, to copy into the deck.

Standalone one-slide .pptx: open it, click the table border, copy, paste into
the deck. Pasting a real table (rather than an image) keeps the text editable
and lets PowerPoint re-style it if the theme changes.

Numbers read from the run, not typed.
"""

import os

import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

REPO = '/Users/mariappan.subramanian/Documents/repo/meridian/demo/synthetic'
RUN = os.path.join(REPO, 'fitted_models', 'scratch_ablation_r90_wellspec')
HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, 'mroi_table.pptx')

INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x5A, 0x5A, 0x5A)
RED = RGBColor(0xC0, 0x39, 0x2B)
TINT = RGBColor(0xEA, 0xF2, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ROW_ALT = RGBColor(0xF7, 0xF9, 0xFB)

# Stops at 5x: 10x is not a spend level anyone plans against, and quoting it
# invites the objection that the test was rigged by extrapolating absurdly.
# The argument is already won at 2-3x.
MULTS = [1.0, 2.0, 3.0, 5.0]
# Both columns are the DEFAULT prior. The contrast between channels IS the
# point: the same prior, the same extrapolation, one channel fine.
CHANNELS = [('TV', 'Channel-1'), ('Display', 'Channel-2')]


def rows():
  m = pd.read_csv(os.path.join(RUN, 'mroi_both_channels.csv'))
  fit = m[m.variant == 'default']
  out = []
  for mult in MULTS:
    cells = [f'{int(mult)}x']
    for key, _ in CHANNELS:
      s = fit[(fit.channel == key) & (fit.spend_multiplier == mult)]
      cells.append(f'{s["pct_error"].median():+.1f}%'.replace('+', ''))
    out.append(tuple(cells))
  return out


def main():
  prs = Presentation()
  prs.slide_width = Inches(13.333)
  prs.slide_height = Inches(7.5)
  slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank

  data = rows()
  headers = ['Spend'] + [f'{label} median error' for _, label in CHANNELS]
  n_rows, n_cols = len(data) + 1, len(headers)

  shape = slide.shapes.add_table(
      n_rows, n_cols, Inches(0.6), Inches(0.6), Inches(6.5),
      Inches(0.42 * n_rows))
  table = shape.table
  table.columns[0].width = Inches(1.3)
  table.columns[1].width = Inches(2.6)
  table.columns[2].width = Inches(2.6)

  # python-pptx applies a banded theme style by default; every fill is set
  # explicitly below so the pasted table does not inherit stripes that fight
  # the deck.
  for c, text in enumerate(headers):
    cell = table.cell(0, c)
    cell.fill.solid()
    cell.fill.fore_color.rgb = TINT
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = cell.margin_right = Inches(0.12)
    p = cell.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = text
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = INK
    r.font.name = 'Calibri'
    if c:
      p.alignment = PP_ALIGN.RIGHT

  for i, row in enumerate(data, start=1):
    for c, text in enumerate(row):
      cell = table.cell(i, c)
      cell.fill.solid()
      cell.fill.fore_color.rgb = WHITE if i % 2 else ROW_ALT
      cell.vertical_anchor = MSO_ANCHOR.MIDDLE
      cell.margin_left = cell.margin_right = Inches(0.12)
      p = cell.text_frame.paragraphs[0]
      r = p.add_run()
      r.text = text
      r.font.size = Pt(13)
      r.font.name = 'Calibri'
      # Channel-1 is emphasised in red because that column is the finding;
      # Channel-2 stays neutral so the eye reads it as the control it is.
      r.font.bold = c == 1
      r.font.color.rgb = RED if c == 1 else (INK if c == 0 else MUTED)
      if c:
        p.alignment = PP_ALIGN.RIGHT

  for r_ in range(n_rows):
    table.rows[r_].height = Inches(0.42)

  prs.save(OUT)
  print('wrote', OUT)
  print()
  print('tab-separated, for a quick paste:')
  print('\t'.join(headers))
  for row in data:
    print('\t'.join(row))


if __name__ == '__main__':
  main()
