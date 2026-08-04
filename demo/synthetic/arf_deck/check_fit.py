"""Text-fit QA without a renderer.

The PowerPoint export path started failing with sandbox error -9074, so this
substitutes for the visual pass on the one defect class that actually matters
here: text overflowing its box. It wraps every run-styled paragraph at the
box's real width using the SAME Calibri files PowerPoint will use, then
compares required height against the box height and against the top of
whatever sits below it.

It cannot see overlap, alignment or colour -- only fit. Stated as such.
"""

import os
import sys

from PIL import ImageFont
from pptx import Presentation
from pptx.util import Pt

FONTS = '/Applications/Microsoft PowerPoint.app/Contents/Resources/DFonts/'
EMU = 914400
DECK = sys.argv[1]
SCALE = 4  # supersample so integer font sizes do not quantise the metrics


def font(pt, bold):
  return ImageFont.truetype(
      FONTS + ('Calibrib.ttf' if bold else 'Calibri.ttf'), int(pt * SCALE))


def wrap(words, width_in, pt, bold):
  """Greedy wrap, returning line count at this width."""
  f = font(pt, bold)
  lines, cur = 1, ''
  for w in words:
    trial = (cur + ' ' + w).strip()
    if f.getlength(trial) / SCALE / 72.0 > width_in and cur:
      lines += 1
      cur = w
    else:
      cur = trial
  return lines


def check(slide_no, slide):
  boxes = []
  for sh in slide.shapes:
    if not sh.has_text_frame or not sh.text_frame.text.strip():
      continue
    w = sh.width / EMU
    h = sh.height / EMU
    top = sh.top / EMU
    need = 0.0
    for para in sh.text_frame.paragraphs:
      runs = [r for r in para.runs if r.text]
      if not runs:
        continue
      # A paragraph can mix sizes/weights; use its largest size and treat the
      # whole paragraph at the widest run's weight, which over-estimates
      # slightly -- the safe direction for an overflow check.
      pt = max((r.font.size.pt if r.font.size else 18) for r in runs)
      bold = any(r.font.bold for r in runs)
      text = ''.join(r.text for r in runs)
      n = wrap(text.split(), w, pt, bold)
      need += n * pt * 1.22 / 72.0  # 1.22 ~ Calibri single line spacing
    left = sh.left / EMU
    boxes.append(
        {'top': top, 'h': h, 'need': need, 'left': left, 'right': left + w,
         'txt': sh.text_frame.text.strip()[:46]})

  boxes.sort(key=lambda b: b['top'])
  print(f'--- slide {slide_no}')
  bad = False
  for i, b in enumerate(boxes):
    status = 'OK  '
    if b['need'] > b['h'] + 0.02:
      status = 'TALL'  # taller than its own box; PowerPoint spills it
    # Only boxes that OVERLAP HORIZONTALLY can collide. Comparing every box
    # against the next one by y alone flags a left-column block against a
    # right-column block sitting at the same height -- which is a two-column
    # layout working correctly, not a defect.
    for other in boxes[i + 1:]:
      overlaps_x = (b['left'] < other['right'] - 0.05
                    and other['left'] < b['right'] - 0.05)
      if overlaps_x and b['top'] + b['need'] > other['top'] + 0.02:
        status = 'CLASH'
        break
    bad |= status != 'OK  '
    print(f'  {status:5s} box h={b["h"]:.2f}" need={b["need"]:.2f}"'
          f'  top={b["top"]:.2f}"  x=[{b["left"]:.2f},{b["right"]:.2f}]'
          f'  | {b["txt"]}')
  return bad


prs = Presentation(DECK)
bad = False
# Every slide, not a hardcoded list: slide indices shift the moment anything
# is inserted, and a fixed list then silently checks the wrong slides.
only = [int(a) for a in sys.argv[2:]] or None
for idx, slide in enumerate(prs.slides):
  if only and idx + 1 not in only:
    continue
  bad |= check(idx + 1, slide)
print('\nRESULT:', 'PROBLEMS FOUND' if bad else 'all text fits')
