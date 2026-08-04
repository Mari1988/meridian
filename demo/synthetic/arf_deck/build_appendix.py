"""Appendix slide: how to override Meridian's three shape priors.

Two lines per parameter — what the default is, and the one line that changes
it. Added as a new slide at the end of the deck.

Every default quoted here is read off `meridian/model/prior_distribution.py`
in the installed source, not from memory.
"""

import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
DECK = os.path.join(HERE, 'ARF-Analytics-Council-Talk-v2.pptx')

INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x5A, 0x5A, 0x5A)
BLUE = RGBColor(0x2C, 0x7F, 0xB8)
RED = RGBColor(0xC0, 0x39, 0x2B)
TINT = RGBColor(0xF3, 0xF6, 0xF9)
MONO = 'Consolas'

# (param, gloss, default-as-shipped, the override line)
PARAMS = [
    (
        'ec_m',
        'Half-saturation — where the curve bends',
        'Default: TruncatedNormal(0.8, 0.8, low=0.1, high=10) — the same '
        'prior for every channel, whatever its audience or reach.',
        'ec_m=tfp.distributions.LogNormal([2.20, 0.26], [0.27, 0.27])',
        'One entry per channel, on the log scale, in multiples of that '
        'channel’s own median execution — log(9.0)=2.20 for Channel-1.',
    ),
    (
        'alpha_m',
        'Adstock decay — how long it carries',
        'Default: Uniform(0, 1) — genuinely flat, and wider than any real '
        'media plan.',
        'alpha_m=tfp.distributions.Uniform([0.3, 0.0], [0.8, 0.3])',
        'A per-medium band is easier to defend than a point estimate: '
        'TV (0.3, 0.8), OOH/print/radio (0.1, 0.4), digital (0, 0.3).',
    ),
    (
        'slope_m',
        'Hill slope — the shape of the bend',
        'Default: Deterministic(1.0) — not a prior at all. It is fixed, so '
        'the model cannot estimate it and never reports uncertainty on it.',
        'slope_m=tfp.distributions.LogNormal([0.7], [0.4])',
        'This is Meridian’s own default for reach/frequency channels '
        '(slope_rf). Overriding it is a real modelling commitment — an '
        'S-shaped response is a stronger claim than a concave one.',
    ),
]

SNIPPET = [
    'from meridian.model import prior_distribution, spec',
    'import tensorflow_probability as tfp',
    '',
    'prior = prior_distribution.PriorDistribution(ec_m=..., alpha_m=..., '
    'slope_m=...)',
    'model_spec = spec.ModelSpec(prior=prior, max_lag=13)',
]


def textbox(slide, x, y, w, h):
  tf = slide.shapes.add_textbox(
      Inches(x), Inches(y), Inches(w), Inches(h)).text_frame
  tf.word_wrap = True
  tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
  return tf


def run(para, text, size, bold=False, color=INK, font='Calibri'):
  r = para.add_run()
  r.text = text
  r.font.size = Pt(size)
  r.font.bold = bold
  r.font.color.rgb = color
  r.font.name = font
  return r


def rule(slide, x, y, w, color, h=0.045):
  bar = slide.shapes.add_shape(
      MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
  bar.fill.solid()
  bar.fill.fore_color.rgb = color
  bar.line.fill.background()
  bar.shadow.inherit = False


def card(slide, x, y, w, h):
  box = slide.shapes.add_shape(
      MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
  box.fill.solid()
  box.fill.fore_color.rgb = TINT
  box.line.fill.background()
  box.shadow.inherit = False
  box.adjustments[0] = 0.08
  box.text_frame.text = ''


def main():
  prs = Presentation(DECK)
  # New slide at the end, on the same Blank layout every other slide uses so
  # it inherits the deck's theme rather than a stock Office one.
  layout = prs.slides[0].slide_layout
  slide = prs.slides.add_slide(layout)
  for sh in list(slide.shapes):
    sh._element.getparent().remove(sh._element)

  tf = textbox(slide, 0.60, 0.35, 12.10, 0.45)
  run(tf.paragraphs[0],
      'Appendix — changing these defaults in Meridian', 24, bold=True)

  tf = textbox(slide, 0.60, 0.88, 12.10, 0.30)
  run(tf.paragraphs[0],
      'All three are priors on ModelSpec. Pass one entry per media channel; '
      'a scalar is broadcast to every channel.', 12, color=MUTED)

  y = 1.42
  for name, gloss, default, code, note in PARAMS:
    rule(slide, 0.60, y, 12.10, BLUE)
    tf = textbox(slide, 0.60, y + 0.13, 12.10, 0.28)
    p = tf.paragraphs[0]
    run(p, name, 14.5, bold=True, color=BLUE, font=MONO)
    run(p, f'  ·  {gloss}', 14.5, bold=True)

    tf = textbox(slide, 0.60, y + 0.46, 12.10, 0.24)
    run(tf.paragraphs[0], default, 11, color=MUTED)

    card(slide, 0.60, y + 0.75, 8.30, 0.34)
    tf = textbox(slide, 0.75, y + 0.81, 8.00, 0.24)
    run(tf.paragraphs[0], code, 11, color=INK, font=MONO)

    tf = textbox(slide, 9.10, y + 0.72, 3.60, 0.60)
    run(tf.paragraphs[0], note, 9.5, color=MUTED)
    y += 1.44

  # max_lag is deliberately outside the three blocks: it is NOT a prior, and
  # the deck's adstock slide makes exactly that point.
  rule(slide, 0.60, y, 12.10, RED)
  tf = textbox(slide, 0.60, y + 0.13, 12.10, 0.28)
  p = tf.paragraphs[0]
  run(p, 'max_lag', 14.5, bold=True, color=RED, font=MONO)
  run(p, '  ·  the carryover window — not a prior, and never estimated',
      14.5, bold=True)
  tf = textbox(slide, 0.60, y + 0.46, 12.10, 0.24)
  run(tf.paragraphs[0],
      'Default 8 weeks, a hard truncation. It is a ModelSpec argument, so no '
      'prior will recover carryover past it: spec.ModelSpec(max_lag=13).',
      11, color=MUTED)

  prs.save(DECK)
  print('added appendix slide ->', DECK, f'({len(prs.slides.__iter__.__self__._sldIdLst)} slides)')


if __name__ == '__main__':
  main()
